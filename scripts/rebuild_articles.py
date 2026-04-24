#!/usr/bin/env python3
"""
assets/articles.md 重建脚本

永远从 Excel 源重建 assets/articles.md，禁止直接编辑 md 文件。
直接编辑 md 会导致标题重复前缀（如"2. 2. 2. 2."）等数据损坏。

用法：
    python3 rebuild_articles.py
"""
from pathlib import Path
from openpyxl import load_workbook

EXCEL_PATH = Path(__file__).resolve().parent.parent / "assets" / "小肥肠文章总结.xlsx"
OUTPUT_PATH = Path(__file__).resolve().parent.parent / "assets" / "articles.md"


def tag_article(a: dict) -> list:
    """给单篇文章打标签"""
    t = a["title"] + " " + a["summary"]
    tags = []

    # 赛道标签
    if "7大赛道" in a["title"]:
        tags.append("入门必看")
    if any(k in t for k in ["小林漫画", "小林"]):
        tags.append("小林漫画")
    if any(k in t for k in ["育儿", "儿童绘本", "绘本"]):
        tags.append("育儿漫画")
    if any(k in t for k in ["趣味漫画", "职场漫画"]):
        tags.append("趣味漫画")
    if any(k in t for k in ["治愈奶奶", "银发", "老奶奶"]):
        tags.append("治愈奶奶")
    if any(k in t for k in ["美食漫画", "美食"]):
        tags.append("美食漫画")
    if any(k in t for k in ["星座漫画", "星座"]):
        tags.append("星座漫画")
    if any(k in t for k in ["老纪", "老纪先生"]):
        tags.append("老纪先生")
    if any(k in t for k in ["小佛陀", "佛陀"]):
        tags.append("小佛陀")

    # 工具标签
    if "Coze" in t:
        tags.append("Coze")
    if "n8n" in t:
        tags.append("n8n")
    if "OpenClaw" in t:
        tags.append("OpenClaw")
    if "liblib" in t:
        tags.append("liblib")
    if any(k in t for k in ["即梦", "明星漫画", "素描"]):
        tags.append("即梦")

    # 内容标签
    if any(k in t for k in ["变现", "商业闭环"]):
        tags.append("变现")
    if any(k in t for k in ["固定IP", "角色"]):
        tags.append("固定IP")

    # 公众号文章标签（排除4篇）
    if "公众号" in a["title"] or "公众号" in a["summary"]:
        excluded_titles = [
            "小红书 3 个虚拟资料赛道",       # 有空格
            "夫妻写公众号年入",              # 年入200万那篇
            "Coze智能体实战：3分钟构建专属数字人！公众号文章一键转为数字人口播视频",
            "最近我的公众号流量几乎腰斩",
        ]
        is_excluded = any(kw in a["title"] for kw in excluded_titles)
        if not is_excluded:
            tags.append("公众号文章")

    return tags


def rebuild():
    if not EXCEL_PATH.exists():
        print(f"错误：找不到 Excel 文件 {EXCEL_PATH}")
        return

    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    articles = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title, link, summary, date = row[0], row[1], row[2] or "", str(row[3]) if row[3] else ""
        if title:
            a = {"title": str(title), "link": link, "summary": summary, "date": date}
            a["tags"] = tag_article(a)
            articles.append(a)

    # 按日期最老排序
    articles.sort(key=lambda x: x["date"])

    # 写入 assets/articles.md
    lines = []
    for i, a in enumerate(articles, 1):
        tag_str = " / ".join(a["tags"]) if a["tags"] else "未分类"
        lines.append(f"### {i}. {a['title']}")
        lines.append(f"- 日期：{a['date']}")
        lines.append(f"- 链接：{a['link']}")
        lines.append(f"- 梗概：{a['summary'][:300]}")
        lines.append(f"- 标签：{tag_str}")
        lines.append("")

    OUTPUT_PATH.write_text("\n".join(lines), encoding='utf-8')

    # 统计
    tag_counts = {}
    for a in articles:
        for t in a["tags"]:
            tag_counts[t] = tag_counts.get(t, 0) + 1

    print(f"重建完成：{len(articles)} 篇")
    print(f"输出文件：{OUTPUT_PATH}")
    print(f"\n标签统计：")
    for tag, count in sorted(tag_counts.items(), key=lambda x: -x[1]):
        print(f"  {tag}: {count}")


if __name__ == "__main__":
    rebuild()
